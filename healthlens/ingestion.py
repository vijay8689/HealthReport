"""Bounded local parsing and conservative source-linked lab extraction."""
import hashlib
import io
import logging
import os
import re
import shutil
import zipfile
from datetime import date, datetime
from pathlib import Path

from healthlens.models import Document, Observation

MAX_BYTES = 20 * 1024 * 1024
MAX_PAGES = 100
MAX_TEXT = 1_000_000
NUMBER = r"[-+]?\d+(?:\.\d+)?"
ROW = re.compile(rf"^\s*(?P<name>[^|]{{1,100}})\|\s*(?P<cmp><=|>=|<|>)?\s*(?P<value>{NUMBER})\s*\|\s*(?P<unit>[^|]{{1,30}})\|\s*(?P<low>{NUMBER})\s*[-–]\s*(?P<high>{NUMBER})\s*$")
BLOCK_NAME = re.compile(r"^[A-Z0-9][A-Z0-9 /,&().\-]{1,100}$")
BLOCK_VALUE = re.compile(r"^\s*(?P<cmp><=|>=|<|>)?\s*(?P<value>[-+]?\d[\d,]*(?:\.\d+)?)\s*$")
BLOCK_RANGE = re.compile(rf"^\s*(?P<low>{NUMBER})\s*[-–]\s*(?P<high>{NUMBER})\s*(?P<unit>[^\d].*)?$")
CATEGORIES = {"hemoglobin": "Hematology", "glucose": "Metabolic", "hba1c": "Metabolic",
              "cholesterol": "Cardiovascular", "triglycer": "Cardiovascular", "vitamin": "Nutritional",
              "creatinine": "Renal", "urea": "Renal", "tsh": "Thyroid"}
NON_TEST_PREFIXES = ("sample collected", "description:", "apollo clinic", "phone no", "email:",
                     "disclaimer:", "lab panel", "health report", "purpose of visit")


class IntakeError(ValueError):
    """Safe, actionable error that can be displayed to the user."""


def tesseract_command() -> str:
    """Resolve OCR on PATH or in standard Windows installer locations."""
    configured = os.getenv("TESSERACT_CMD")
    if configured:
        if Path(configured).is_file():
            return configured
        raise IntakeError("TESSERACT_CMD does not point to a Tesseract executable. Check the configured path.")
    command = shutil.which("tesseract")
    if command:
        return command
    if os.name == "nt":
        candidates = [Path(os.getenv("ProgramFiles", "C:/Program Files")) / "Tesseract-OCR/tesseract.exe",
                      Path(os.getenv("ProgramFiles(x86)", "C:/Program Files (x86)")) / "Tesseract-OCR/tesseract.exe"]
        local = os.getenv("LOCALAPPDATA")
        if local:
            candidates.extend([Path(local) / "Programs/Tesseract-OCR/tesseract.exe",
                               Path(local) / "Tesseract-OCR/tesseract.exe"])
        for candidate in candidates:
            if candidate.is_file():
                return str(candidate)
    raise IntakeError("Tesseract OCR is not installed or could not be found. Install it with English language data, or set TESSERACT_CMD to its executable path.")


def ocr_image(image) -> str:
    try:
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = tesseract_command()
        from PIL import Image, ImageOps
        # Phone images may have rotation metadata, transparency, or palette modes.
        prepared = ImageOps.exif_transpose(image).convert("RGBA")
        background = Image.new("RGBA", prepared.size, "white")
        prepared = Image.alpha_composite(background, prepared).convert("RGB")
        return pytesseract.image_to_string(prepared, lang="eng", timeout=60)
    except ImportError as exc:
        raise IntakeError("OCR is unavailable. Install pytesseract and Tesseract, or upload a text-based report.") from exc
    except IntakeError:
        raise
    except Exception as exc:
        logging.getLogger(__name__).exception("Local OCR failed")
        if isinstance(exc, RuntimeError) and "timeout" in str(exc).lower():
            raise IntakeError("OCR timed out after 60 seconds. Upload a smaller, clear image or split the scanned PDF into smaller files.") from exc
        if isinstance(exc, pytesseract.TesseractNotFoundError):
            raise IntakeError("The Tesseract executable could not start. Check TESSERACT_CMD or reinstall the OCR engine.") from exc
        if isinstance(exc, pytesseract.TesseractError):
            detail = " ".join(str(exc.message).split())[:500]
            raise IntakeError(f"Tesseract OCR failed: {detail}") from exc
        if isinstance(exc, OSError):
            raise IntakeError(f"OCR could not access the engine or temporary files: {exc.strerror or type(exc).__name__}.") from exc
        raise IntakeError(f"OCR could not process this image ({type(exc).__name__}). Try a standard PNG/JPEG image or an unlocked PDF.") from exc


def parse_bytes(content: bytes, filename: str) -> tuple[list[dict], bool]:
    if not content or len(content) > MAX_BYTES:
        raise IntakeError("The document must contain data and be smaller than 20 MB.")
    ext = Path(filename).suffix.lower()
    pages, used_ocr = [], False
    try:
        if ext == ".txt":
            text = content.decode("utf-8-sig")
            if "\x00" in text:
                raise IntakeError("This file is not supported UTF-8 text.")
            pages = [{"locator": "Text", "text": text}]
        elif ext == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise IntakeError("The file signature does not match PDF.")
            import pymupdf
            from PIL import Image
            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                if pdf.needs_pass:
                    raise IntakeError("Password-protected PDFs are not supported. Upload an unlocked copy.")
                if len(pdf) > MAX_PAGES:
                    raise IntakeError("This report exceeds the 100-page limit.")
                scanned = 0
                for index, page in enumerate(pdf):
                    text = page.get_text()
                    page_ocr = len(text.strip()) < 15
                    if page_ocr:
                        scanned += 1
                        if scanned > 5:
                            raise IntakeError("Demo OCR supports up to five scanned pages per file. Split this report.")
                        if page.rect.width * page.rect.height * 4 > 20_000_000:
                            raise IntakeError("This PDF page exceeds the safe raster size.")
                        pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
                        text = ocr_image(Image.open(io.BytesIO(pix.tobytes("png"))))
                        used_ocr = True
                    record = {"locator": f"Page {index + 1}", "text": text}
                    if not page_ocr:
                        record["lab_rows"] = _pdf_lab_rows(page, record["locator"])
                    pages.append(record)
        elif ext == ".docx":
            if not zipfile.is_zipfile(io.BytesIO(content)):
                raise IntakeError("The file signature does not match DOCX.")
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                    raise IntakeError("The expanded document exceeds the safe size limit.")
                if "word/document.xml" not in archive.namelist():
                    raise IntakeError("This archive is not a Word document.")
            from docx import Document as WordDocument
            doc = WordDocument(io.BytesIO(content))
            pages = [{"locator": f"Paragraph {i + 1}", "text": p.text}
                     for i, p in enumerate(doc.paragraphs) if p.text.strip()]
            for ti, table in enumerate(doc.tables):
                for ri, row in enumerate(table.rows):
                    pages.append({"locator": f"Table {ti + 1}, row {ri + 1}",
                                  "text": " | ".join(c.text for c in row.cells)})
        elif ext in {".png", ".jpg", ".jpeg"}:
            from PIL import Image
            with Image.open(io.BytesIO(content)) as image:
                if image.format not in {"PNG", "JPEG"} or image.width * image.height > 20_000_000:
                    raise IntakeError("Use a PNG/JPEG image under 20 megapixels.")
                pages = [{"locator": "Image 1", "text": ocr_image(image)}]
                used_ocr = True
        else:
            raise IntakeError("Use PDF, DOCX, TXT, PNG, JPG, or JPEG.")
    except IntakeError:
        raise
    except Exception as exc:
        raise IntakeError("The document could not be read. Check its format and try an unlocked, undamaged copy.") from exc
    if sum(len(page["text"]) for page in pages) > MAX_TEXT:
        raise IntakeError("The extracted text exceeds the demo processing limit.")
    if not any(page["text"].strip() for page in pages):
        raise IntakeError("No readable text was found in this document.")
    return pages, used_ocr


def normalize_patient_name(value: str) -> str:
    value = re.sub(r"^(?:(?:mrs|mr|ms|miss|dr)(?:\.\s*|\s+))+", "", value.strip(), flags=re.I)
    return " ".join(value.casefold().split())


def patient_name_matches(profile_name: str, report_name: str) -> bool:
    profile = normalize_patient_name(profile_name).split()
    report = normalize_patient_name(report_name).split()
    # A profile may use a first name or first/middle names from the full report name.
    return bool(profile) and report[:len(profile)] == profile


def _patient_names(pages: list[dict]) -> set[str]:
    text = "\n".join(page["text"] for page in pages)
    pattern = r"^[ \t]*(?:Patient[ \t]+Name|Patient|Name)(?:[ \t]*\([^\r\n)]*\))?\s*:[ \t]*(?:\r?\n[ \t]*)?([^\r\n]+)"
    names = set()
    for match in re.finditer(pattern, text, re.I | re.M):
        parts = [match.group(1).strip()]
        # Some PDF headers wrap the patient name across consecutive text lines.
        for line in text[match.end():].splitlines()[1:]:
            line = line.strip()
            if not line or not re.fullmatch(r"[^\W\d_]+(?:[ .'-]+[^\W\d_]+)*\.?", line, re.UNICODE):
                break
            if re.match(r"(?:patient|age|gender|sex|date|report|sample|test|health|doctor|referr)\b", line, re.I):
                break
            parts.append(line)
        names.add(normalize_patient_name(" ".join(parts)))
    return names



def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\ufffd", " ").replace("Â", " ")).strip()


def _looks_like_name(value: str) -> bool:
    if not value or len(value) > 110 or value.lower().startswith(NON_TEST_PREFIXES):
        return False
    letters = [char for char in value if char.isalpha()]
    return bool(letters) and BLOCK_NAME.match(value) is not None and sum(c.isupper() for c in letters) / len(letters) > .72


def _observation_dates(text: str) -> set[date]:
    # Collection dates describe the measurements; receipt/print dates do not.
    labels = [r"(?:Sample\s+Collected\s+on|Collected|Collection\s+Date|Sample\s+Collection\s+Date)",
              r"(?:Observation\s+Date|Date)", r"(?:Reported|Report\s+Date)"]
    token = r"(\d{4}-\d{2}-\d{2}|\d{1,2}[-/](?:[A-Za-z]{3,9}|\d{1,2})[-/]\d{4})"
    for label in labels:
        values = re.findall(rf"^\s*{label}\s*:\s*{token}", text, re.I | re.M)
        if not values:
            continue
        dates = set()
        for value in values:
            for fmt in ("%Y-%m-%d", "%d/%b/%Y", "%d-%b-%Y", "%d/%B/%Y", "%d-%B-%Y", "%d/%m/%Y", "%d-%m-%Y"):
                try:
                    dates.add(datetime.strptime(value, fmt).date())
                    break
                except ValueError:
                    continue
        return dates
    return set()


def _page_date(text: str, fallback: date | None) -> date | None:
    dates = _observation_dates(text)
    return next(iter(dates)) if len(dates) == 1 else (None if dates else fallback)


def _pdf_lab_rows(page, locator: str) -> list[dict]:
    """Read ruled lab tables with an explicit five-column source header."""
    tables = page.find_tables().tables
    headers = []
    for table in tables:
        for cells in table.extract():
            if [_clean(cell or "").casefold() for cell in cells] == ["test name", "result", "unit", "bio. ref. range", "method"]:
                headers.append(table.bbox)
    if not headers:
        return []
    rows = []
    for table_index, table in enumerate(tables, 1):
        if table.col_count != 5 or not any(
            (table.bbox == header or table.bbox[1] >= header[3])
            and abs(table.bbox[0] - header[0]) < 3 and abs(table.bbox[2] - header[2]) < 3
            for header in headers
        ):
            continue
        for row_index, cells in enumerate(table.extract(), 1):
            if len(cells) != 5 or not cells[0] or not cells[1] or _clean(cells[0]).casefold() == "test name":
                continue
            rows.append(dict(cells=[cell or "" for cell in cells],
                             locator=f"{locator}, table {table_index}, row {row_index}"))
    return rows


def _table_observations(page: dict, patient_id: str, document_id: str,
                        fallback_date: date | None) -> list[Observation]:
    observations = []
    for row in page.get("lab_rows", []):
        name, value, unit, reference, method = [_clean(cell) for cell in row["cells"]]
        number = BLOCK_VALUE.fullmatch(value)
        if not number:
            continue  # Qualitative results, titres and intervals remain source evidence.
        low = high = None
        if reference:
            limits = re.fullmatch(rf"({NUMBER})\s*[-\u2013\u2014]\s*({NUMBER})\.?", reference)
            if limits:
                low, high = map(float, limits.groups())
        try:
            observations.append(Observation(
                patient_id=patient_id, document_id=document_id, name=name,
                category=next((category for token, category in CATEGORIES.items() if token in name.lower()), "General"),
                value=float(number.group("value").replace(",", "")), original_value=value,
                comparator=number.group("cmp") or "=", unit=unit or "Not recorded",
                low=low, high=high, date=_page_date(page["text"], fallback_date),
                method=method or "Not recorded", locator=row["locator"],
                source_text="\n".join(cell for cell in row["cells"] if cell),
            ))
        except ValueError:
            continue
    return observations


def _block_observations(page: dict, patient_id: str, document_id: str,
                        fallback_date: date | None) -> list[Observation]:
    raw = page["text"].splitlines()
    lines = [_clean(line) for line in raw]
    observed_on = _page_date(page["text"], fallback_date)
    results = []
    for index in range(2, len(lines)):
        range_match = BLOCK_RANGE.match(lines[index])
        value_match = BLOCK_VALUE.match(lines[index - 1])
        if not (range_match and value_match):
            continue
        names, cursor = [], index - 2
        while cursor >= 0 and len(names) < 3 and _looks_like_name(lines[cursor]):
            names.insert(0, lines[cursor])
            cursor -= 1
        if not names:
            continue
        name = " ".join(names)
        unit = _clean(range_match.group("unit") or "").strip(" *") or "Not recorded"
        value_text = value_match.group("value")
        locator_start = index - len(names)
        category = next((category for token, category in CATEGORIES.items() if token in name.lower()), "General")
        try:
            results.append(Observation(
                patient_id=patient_id, document_id=document_id, name=name, category=category,
                value=float(value_text.replace(",", "")), original_value=(value_match.group("cmp") or "") + value_text,
                comparator=value_match.group("cmp") or "=", unit=unit,
                low=float(range_match.group("low")), high=float(range_match.group("high")), date=observed_on,
                locator=f"{page['locator']}, lines {locator_start + 1}-{index + 1}",
                source_text="\n".join(raw[locator_start:index + 1]),
            ))
        except ValueError:
            continue
    return results


def ingest(content: bytes, filename: str, patient_id: str,
           patient_name: str | None = None) -> tuple[Document, list[Observation]]:
    pages, used_ocr = parse_bytes(content, filename)
    text = "\n".join(page["text"] for page in pages)
    names = _patient_names(pages)
    target_patient_id = patient_id
    mismatch = patient_name is not None and (
        len(names) != 1 or not patient_name_matches(patient_name, next(iter(names)))
    )
    dates = set().union(*(_observation_dates(page["text"]) for page in pages))
    report_date = next(iter(dates)) if len(dates) == 1 else None
    doc = Document(patient_id=target_patient_id, name=Path(filename).name, date=report_date,
                   hash=hashlib.sha256(content).hexdigest(), pages=pages, ocr=used_ocr)
    if mismatch or not names:
        doc.status = "Quarantined"
        if not names:
            doc.warnings.append("No explicit patient name was found. Check the patient-name header in the source document.")
        else:
            detected = "; ".join(sorted(names))
            expected = patient_name or "Not specified"
            doc.warnings.append(
                f"Patient name mismatch. Selected profile: {expected}. Detected report name(s): {detected}. "
                "Select the matching patient profile or check the extracted source text. Patient IDs are not checked."
            )
        return doc, []
    if not report_date:
        doc.warnings.append("No single observation date could be determined for this document. Results keep any collection dates found on their source pages.")
    observations = []
    for page in pages:
        if page.get("lab_rows"):
            observations.extend(_table_observations(page, target_patient_id, doc.id, report_date))
            continue
        for line_no, line in enumerate(page["text"].splitlines(), 1):
            match = ROW.match(line)
            if not match:
                continue
            row = match.groupdict()
            name = row["name"].strip()
            category = next((value for key, value in CATEGORIES.items() if key in name.lower()), "General")
            locator = f"Line {line_no}" if page["locator"] == "Text" else f"{page['locator']}, line {line_no}"
            try:
                observations.append(Observation(
                    patient_id=target_patient_id, document_id=doc.id, name=name, category=category,
                    value=float(row["value"]), original_value=(row["cmp"] or "") + row["value"],
                    comparator=row["cmp"] or "=", unit=row["unit"].strip(), low=float(row["low"]),
                    high=float(row["high"]), date=report_date, locator=locator, source_text=line,
                ))
            except ValueError:
                doc.warnings.append(f"Invalid range/value at {locator}; row excluded.")
        observations.extend(_block_observations(page, target_patient_id, doc.id, report_date))
    unique = {}
    for observation in observations:
        key = (observation.name, observation.date, observation.value,
               observation.unit, observation.low, observation.high, observation.method, observation.specimen)
        unique.setdefault(key, observation)
    observations = list(unique.values())
    if not observations:
        doc.warnings.append("No supported numeric laboratory rows found. Source text remains available in Evidence and Documents.")
        doc.status = "Text only"
    elif any(o.low == 0 and o.high == 0 for o in observations):
        doc.warnings.append("The source contains one or more 0-0 reference ranges. Review them before use; they are not treated as validated thresholds.")
    non_numeric = sum(not BLOCK_VALUE.fullmatch(_clean(row["cells"][1]))
                      for page in pages for row in page.get("lab_rows", []))
    if non_numeric:
        doc.warnings.append(f"{non_numeric} text, titre, or interval results are retained in source evidence and excluded from numeric trends.")
    for observation in observations:
        observation.status = "accepted"
    if observations:
        doc.status = "Imported"
    return doc, observations
